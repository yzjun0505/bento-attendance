from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


OUT = Path("/Users/eva/Desktop/Project/境图考勤系统需求分析说明书.docx")
ASSET_DIR = Path("/Users/eva/Desktop/Project/.doc_assets")
ASSET_DIR.mkdir(exist_ok=True)

FONT_PATH = Path("/System/Library/Fonts/STHeiti Medium.ttc")
FONT_CN = str(FONT_PATH if FONT_PATH.exists() else "/System/Library/Fonts/Supplemental/Arial Unicode.ttf")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9E2EC", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths_cm):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)


def style_table(table, header=True):
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "宋体"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                    run.font.size = Pt(9)
            if header and r_idx == 0:
                set_cell_shading(cell, "EEF2F7")
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(31, 77, 120)


def add_kv_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_width(table, [4.2, 10.6])
    for i, (key, value) in enumerate(rows):
        table.cell(i, 0).text = key
        table.cell(i, 1).text = value
        set_cell_shading(table.cell(i, 0), "EEF2F7")
    style_table(table, header=False)
    return table


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    style_table(table)
    return table


def add_para(doc, text="", style=None, align=None):
    p = doc.add_paragraph(text, style=style)
    if align:
        p.alignment = align
    return p


def add_bullets(doc, items):
    for item in items:
        add_para(doc, item, "List Bullet")


def add_numbered(doc, items):
    for item in items:
        add_para(doc, item, "List Number")


def font(size=28, bold=False):
    return ImageFont.truetype(FONT_CN, size=size, index=0)


def rounded_box(draw, xy, text, fill, outline="#4A5568", w=220, h=70, size=24):
    x, y = xy
    draw.rounded_rectangle([x, y, x + w, y + h], radius=14, fill=fill, outline=outline, width=2)
    lines = wrap(text, 10)
    total = len(lines) * (size + 4)
    for idx, line in enumerate(lines):
        tw = draw.textlength(line, font=font(size))
        draw.text((x + (w - tw) / 2, y + (h - total) / 2 + idx * (size + 4)), line, fill="#1F2937", font=font(size))


def arrow(draw, start, end):
    draw.line([start, end], fill="#4A5568", width=3)
    ex, ey = end
    sx, sy = start
    if ex >= sx:
        pts = [(ex, ey), (ex - 12, ey - 7), (ex - 12, ey + 7)]
    else:
        pts = [(ex, ey), (ex + 12, ey - 7), (ex + 12, ey + 7)]
    draw.polygon(pts, fill="#4A5568")


def make_architecture():
    path = ASSET_DIR / "architecture.png"
    img = Image.new("RGB", (1400, 820), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "境图考勤系统总体架构", fill="#111827", font=font(34))
    rounded_box(d, (70, 150), "Flutter移动端\n打卡/定位/审批/IM", "#E8F4FF", w=300, h=120, size=24)
    rounded_box(d, (70, 360), "Vue管理端\n人员/排班/报表", "#F0FDF4", w=300, h=120, size=24)
    rounded_box(d, (530, 250), "Node.js Express API\n认证/业务/文件/AI", "#FFF7ED", w=330, h=150, size=24)
    rounded_box(d, (1030, 120), "MySQL\n业务数据", "#F8FAFC", w=260, h=90, size=24)
    rounded_box(d, (1030, 260), "MongoDB\nAI/扩展数据", "#F8FAFC", w=260, h=90, size=24)
    rounded_box(d, (1030, 400), "OpenIM\n即时通讯", "#F8FAFC", w=260, h=90, size=24)
    rounded_box(d, (1030, 540), "高德地图/推送\n位置与通知服务", "#F8FAFC", w=260, h=90, size=24)
    arrow(d, (370, 210), (530, 300))
    arrow(d, (370, 420), (530, 355))
    for y in [165, 305, 445, 585]:
        arrow(d, (860, 325), (1030, y))
    d.text((85, 650), "说明：客户端通过 HTTPS/局域网 API 访问后端；后端统一完成权限校验、数据持久化、第三方服务调用与审计记录。", fill="#374151", font=font(23))
    img.save(path)
    return path


def make_use_case():
    path = ASSET_DIR / "use_case.png"
    img = Image.new("RGB", (1400, 900), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "系统用例关系图", fill="#111827", font=font(34))
    actors = [("员工", 90, 190), ("管理员/主管", 90, 610)]
    for name, x, y in actors:
        d.ellipse([x + 60, y, x + 110, y + 50], outline="#374151", width=3)
        d.line([x + 85, y + 50, x + 85, y + 140], fill="#374151", width=3)
        d.line([x + 30, y + 85, x + 140, y + 85], fill="#374151", width=3)
        d.line([x + 85, y + 140, x + 35, y + 210], fill="#374151", width=3)
        d.line([x + 85, y + 140, x + 135, y + 210], fill="#374151", width=3)
        d.text((x + 35, y + 225), name, fill="#111827", font=font(25))
    use_cases = [
        ("登录与会话管理", 420, 120), ("外勤打卡", 420, 250), ("位置上报与轨迹查看", 420, 380),
        ("提交审批申请", 420, 510), ("消息与AI助手", 420, 640), ("人员/项目管理", 880, 160),
        ("排班与节假日配置", 880, 300), ("考勤统计与异常处理", 880, 440),
        ("水印模板与防伪查询", 880, 580), ("通知发送与报表导出", 880, 720),
    ]
    for text, x, y in use_cases:
        d.ellipse([x, y, x + 310, y + 78], fill="#F8FAFC", outline="#64748B", width=2)
        tw = d.textlength(text, font=font(23))
        d.text((x + (310 - tw) / 2, y + 25), text, fill="#1F2937", font=font(23))
    for y in [159, 289, 419, 549, 679]:
        d.line([(225, 300), (420, y)], fill="#9CA3AF", width=2)
    for y in [199, 339, 479, 619, 759]:
        d.line([(225, 720), (880, y)], fill="#9CA3AF", width=2)
    img.save(path)
    return path


def make_flow():
    path = ASSET_DIR / "checkin_flow.png"
    img = Image.new("RGB", (1400, 620), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "外勤打卡业务流程", fill="#111827", font=font(34))
    steps = ["进入打卡页", "获取定位\n与项目围栏", "拍照并叠加\n时间地点水印", "提交打卡\n或离线缓存", "后端校验\n距离/权限/类型", "生成记录\n统计与通知"]
    xs = [55, 270, 495, 720, 945, 1170]
    for x, text in zip(xs, steps):
        rounded_box(d, (x, 210), text, "#EEF2FF", w=175, h=105, size=21)
    for x in xs[:-1]:
        arrow(d, (x + 175, 262), (x + 215, 262))
    d.text((80, 430), "异常路径：定位失败时提示重试；网络不可用时写入本地离线队列；围栏外打卡记录 is_outside 与距离，供主管复核。", fill="#374151", font=font(24))
    img.save(path)
    return path


def make_er():
    path = ASSET_DIR / "er.png"
    img = Image.new("RGB", (1400, 860), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "核心数据实体关系图", fill="#111827", font=font(34))
    boxes = {
        "users\n用户": (80, 130), "projects\n项目": (360, 130), "checkins\n打卡记录": (640, 130), "locations\n位置记录": (940, 130),
        "attendance_groups\n考勤组": (80, 420), "user_schedules\n排班": (360, 420), "approval_requests\n审批申请": (640, 420), "notifications\n通知": (940, 420),
        "watermark_templates\n水印模板": (220, 650), "ai_conversations\nAI会话": (520, 650), "sessions\n会话令牌": (820, 650),
    }
    for name, (x, y) in boxes.items():
        rounded_box(d, (x, y), name, "#F8FAFC", w=220, h=92, size=21)
    links = [
        ((300, 176), (360, 176)), ((580, 176), (640, 176)), ((300, 176), (940, 176)),
        ((190, 222), (190, 420)), ((470, 222), (470, 420)), ((750, 222), (750, 420)),
        ((1050, 222), (1050, 420)), ((190, 512), (330, 650)), ((750, 512), (630, 650)), ((1050, 512), (930, 650)),
    ]
    for a, b in links:
        d.line([a, b], fill="#64748B", width=3)
    d.text((80, 790), "注：图中仅展示核心实体；设备、班次、节假日、离线打卡、AI操作审计等表在逻辑结构中单独说明。", fill="#374151", font=font(23))
    img.save(path)
    return path


def set_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color, before, after in [
        ("Title", 24, "000000", 0, 10),
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "黑体" if name != "Normal" else "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for list_style in ["List Bullet", "List Number"]:
        style = styles[list_style]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    set_styles(doc)

    p = add_para(doc, "医学信息工程学院", align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in p.runs:
        r.font.size = Pt(16)
        r.bold = True
    p = add_para(doc, "《软件工程课程实践》", align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in p.runs:
        r.font.size = Pt(18)
        r.bold = True
    doc.add_paragraph()
    p = add_para(doc, "需 求 分 析", align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in p.runs:
        r.font.size = Pt(28)
        r.bold = True
        r.font.name = "黑体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    doc.add_paragraph()
    add_kv_table(doc, [
        ("项目名称", "境图考勤系统"),
        ("姓名", "（请填写）"),
        ("部系", "医学信息工程学院"),
        ("专业", "（请填写）"),
        ("年级", "（请填写）"),
        ("Email", "（请填写）"),
        ("联系电话", "（请填写）"),
        ("指导教师", "（请填写）"),
    ])
    doc.add_page_break()

    doc.add_heading("一、绪论", level=1)
    doc.add_heading("1.1 引言", level=2)
    add_para(doc, "随着移动互联网、地理信息服务和企业数字化管理的快速发展，传统依赖纸质签到、人工汇总和事后核对的考勤方式已经难以满足外勤、项目制和多地点办公场景的管理需求。员工在项目现场、巡检途中或临时办公点完成工作时，管理者不仅需要知道是否按时出勤，还需要掌握打卡地点、照片凭证、轨迹记录、异常原因和审批处理过程。")
    add_para(doc, "境图考勤系统是一套面向企业外勤与项目现场管理的多端协同考勤平台，系统由 Flutter 移动端、Vue 管理后台、Node.js 后端服务以及 MySQL、MongoDB、OpenIM 等基础设施组成。系统围绕“真实位置、可信凭证、及时协同、可追溯分析”的目标，提供外勤打卡、实时定位、水印防伪、离线同步、排班管理、审批流、消息通知、轨迹回放和 AI 数据助理等功能。")
    doc.add_heading("1.2 课题研究的意义", level=2)
    add_para(doc, "本课题将软件工程课程中的需求分析、系统建模、数据库设计、接口设计、测试验证等方法应用于实际业务系统，具有较强的实践价值。其意义主要体现在以下几个方面：")
    add_numbered(doc, [
        "从企业管理角度看，系统能够降低考勤统计和现场核验的人力成本，减少口头汇报、纸质签到和手工汇总造成的数据滞后。",
        "从员工使用角度看，移动端支持拍照打卡、离线缓存、审批申请和消息查看，能够适应网络不稳定或外勤地点复杂的真实工作环境。",
        "从数据可信角度看，系统通过定位围栏、水印防伪码、照片凭证、轨迹记录和操作审计形成可追溯链路，提升考勤数据的真实性。",
        "从技术实践角度看，项目覆盖移动端、Web 管理端、RESTful API、关系型数据库、即时通讯、地图服务和 AI 工具调用等技术，能够体现完整的软件工程开发过程。"
    ])
    doc.add_heading("1.3 系统设计的内容", level=2)
    add_para(doc, "境图考勤系统的核心工作流程为：员工登录移动端后，根据排班和项目位置进入打卡页面；移动端获取定位、项目围栏和水印模板，员工拍照后生成带时间、地点、人员和防伪码的照片凭证；若网络可用则立即提交至后端，若网络不可用则写入本地离线队列并在恢复网络后自动同步。管理端提供人员、项目、设备、排班、节假日、打卡记录、轨迹回放、审批和水印模板等后台能力，主管可对异常打卡、请假补卡、外勤轨迹和统计报表进行管理。")
    add_para(doc, "系统还集成 OpenIM 即时通讯和 AI 数据助理。AI 助手通过后端受控工具查询考勤、项目、审批和报表数据，并对导出报表、发送通知、创建排班等高风险操作采用确认式执行机制，避免模型直接越权修改业务数据。")

    doc.add_heading("二、系统描述", level=1)
    doc.add_heading("2.1 系统设计目标", level=2)
    add_bullets(doc, [
        "功能完整性：覆盖员工移动打卡、定位轨迹、离线同步、审批申请、消息通知、排班统计和后台管理等主要流程。",
        "业务可信性：通过围栏距离、照片水印、防伪码、位置记录和审计日志提升考勤数据可信度。",
        "多端一致性：移动端、管理端和后端接口围绕统一数据模型协作，保证考勤记录、排班规则和审批状态一致。",
        "可维护性：采用分层后端结构，路由、控制器、服务、模型职责清晰，前端按页面和 API 模块组织。",
        "可扩展性：支持新增打卡类型、项目围栏、AI 工具、通知渠道和第三方地图/通讯服务。"
    ])
    doc.add_heading("2.2 系统体系结构", level=2)
    add_kv_table(doc, [
        ("系统架构", "移动端 + Web 管理端 + RESTful 后端服务 + 数据库/第三方服务"),
        ("移动端技术", "Flutter、Dio、BLoC、SharedPreferences、Camera、Geolocator、WebView、高德地图"),
        ("管理端技术", "Vue 3、Vite、Element Plus、Pinia、Vue Router、Axios、ECharts"),
        ("后端技术", "Node.js、Express、JWT、Multer、ExcelJS、Socket.IO、node-cron、Winston"),
        ("数据库", "MySQL 存储核心业务数据，MongoDB 用于可选扩展与 AI 相关数据"),
        ("第三方服务", "OpenIM 即时通讯、高德地图定位与逆地理编码、移动端本地/推送通知服务"),
        ("开发与测试", "Jest、Supertest、Flutter Test、Docker Compose、RESTful API 文档")
    ])
    arch = make_architecture()
    doc.add_picture(str(arch), width=Cm(15.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("2.3 系统使用技术", level=2)
    doc.add_heading("2.3.1 安全防护体系", level=3)
    add_para(doc, "系统采用多层安全防护机制，围绕身份认证、接口访问、数据隔离和操作审计进行设计。后端通过 JWT 令牌识别当前用户身份，移动端和管理端在请求业务接口时携带认证信息；用户密码使用 bcryptjs 哈希存储；刷新令牌和设备信息保存在 sessions 表中，便于会话续期和失效控制。")
    add_bullets(doc, [
        "认证安全：登录成功后签发访问令牌，敏感接口由 auth 中间件统一校验。",
        "角色控制：系统区分 admin、manager、worker 等角色，管理端能力与当前登录用户权限绑定。",
        "接口防护：后端使用 express-rate-limit、CORS、压缩和统一异常处理降低常见接口风险。",
        "数据隔离：员工只能查看和提交与自己相关的移动端数据，管理端按业务角色处理人员、项目和考勤数据。",
        "审计留痕：AI 操作、审批处理、打卡记录、防伪码状态和会话信息均可形成追溯依据。"
    ])
    doc.add_heading("2.3.2 定位围栏与距离判定", level=3)
    add_para(doc, "外勤打卡的核心约束是“人在指定项目地点附近完成可信打卡”。系统在 projects 表中维护项目地址、经纬度和打卡半径，移动端通过高德地图获取当前定位并展示项目位置，后端保存打卡点经纬度、地址、是否围栏外打卡以及 distance_to_fence 距离字段。")
    add_numbered(doc, [
        "移动端获取当前位置和精度信息，必要时通过高德服务进行逆地理编码。",
        "系统读取当前员工绑定项目或选择的项目，得到项目中心点与允许半径。",
        "根据两点经纬度计算距离，若距离大于半径则标记为围栏外打卡。",
        "围栏外打卡不直接丢弃，而是记录异常原因，供主管在管理端复核。"
    ])
    doc.add_heading("2.3.3 水印防伪机制", level=3)
    add_para(doc, "系统在打卡拍照环节自动叠加时间、地点、人员、项目和防伪码等信息。防伪码由后端生成并维护有效期，打卡记录写入 watermark_code 字段，管理端提供防伪码查询入口，能够验证照片对应的打卡记录、人员、地点和时间。")
    add_bullets(doc, [
        "水印模板可在管理端维护，移动端根据模板 schema_json 渲染样式。",
        "防伪码与用户、创建时间、过期时间、使用状态关联，防止重复伪造。",
        "照片路径与打卡记录绑定，后续可结合防伪码查询和打卡详情完成核验。"
    ])
    doc.add_heading("2.3.4 离线同步策略", level=3)
    add_para(doc, "针对外勤场景中经常出现的弱网或无网情况，移动端实现离线打卡缓存。无网络时，打卡类型、位置、照片、备注、项目和本地时间会先写入本地队列；恢复网络后，移动端调用 /api/offline-checkins/sync 接口提交缓存数据，后端将其转为正式打卡记录并标记同步状态。")
    add_bullets(doc, [
        "本地缓存保证员工在无网场景下仍可完成打卡动作。",
        "local_timestamp 保存设备本地打卡时间，避免同步时间替代真实发生时间。",
        "synced 与 sync_time 字段用于区分待同步、已同步和同步失败记录。"
    ])
    doc.add_heading("2.3.5 考勤分析与排班规则", level=3)
    add_para(doc, "后端的考勤分析服务通过 node-cron 定时执行，按考勤组、排班、班次时间和打卡记录生成 attendance_results。系统支持迟到容忍、早退容忍、休息日、节假日、请假审批联动等业务规则，为工作台统计、异常打卡和报表导出提供基础数据。")
    doc.add_heading("2.3.6 AI 数据助理与受控工具调用", level=3)
    add_para(doc, "AI 助手并不直接操作数据库，而是通过后端定义的受控工具读取或处理业务数据。查询类能力可以返回统计卡片、表格和图表结构；执行类能力需要生成待确认操作，由用户二次确认后再落地，确认记录写入 ai_action_requests 和 ai_audit_logs。该设计能够兼顾智能交互与业务安全。")

    doc.add_heading("三、可行性分析", level=1)
    doc.add_heading("3.1 经济可行性", level=2)
    add_para(doc, "本系统采用开源技术栈实现，Flutter、Vue、Node.js、MySQL、MongoDB、Jest 等工具均可低成本使用。对于课程实践或中小规模企业内部系统，部署可先基于本地服务器或云主机完成，随着数据规模增长再扩展数据库、对象存储和消息服务。因此系统在开发成本、运行成本和后续维护成本方面具有可行性。")
    doc.add_heading("3.2 操作可行性", level=2)
    add_para(doc, "移动端面向普通员工，核心操作集中在登录、打卡、拍照、查看记录、提交审批和查看消息；管理端面向主管和管理员，采用侧边栏导航和表格表单模式组织人员、项目、排班、审批和报表。整体交互符合常见移动办公和后台管理系统习惯，用户学习成本较低。")
    doc.add_heading("3.3 技术可行性", level=2)
    add_para(doc, "项目中的关键技术均有成熟生态支持。Flutter 可同时覆盖 Android、iOS 和桌面调试场景；Vue 3 与 Element Plus 能快速构建管理端；Express 适合 RESTful API 与中间件扩展；MySQL 适合结构化业务数据；OpenIM、高德地图和推送服务可作为外部能力接入。因此系统在技术选型和工程落地上具有可行性。")

    doc.add_heading("四、需求分析", level=1)
    doc.add_heading("4.1 系统功能模块的划分", level=2)
    add_table(doc, ["模块", "主要使用者", "功能范围"], [
        ["用户与认证管理", "员工、管理员", "注册/登录、会话续期、个人资料、角色权限、密码修改"],
        ["移动打卡管理", "员工", "项目打卡、拍照水印、围栏判断、离线缓存、自定义打卡类型"],
        ["定位与轨迹管理", "员工、主管", "实时位置上报、轨迹回放、热力图、停留点分析"],
        ["排班与考勤统计", "主管、管理员", "班次、排班、节假日、考勤组、迟到早退缺勤统计"],
        ["审批管理", "员工、主管", "请假、补卡、加班申请与审批处理"],
        ["项目与设备管理", "管理员", "项目地点、围栏半径、设备绑定与状态维护"],
        ["消息与即时通讯", "全体用户", "系统通知、OpenIM 聊天、未读提醒"],
        ["水印与防伪查询", "员工、主管", "水印模板、打卡防伪码、照片凭证核验"],
        ["AI 数据助理", "主管、管理员", "自然语言查数据、结构化展示、确认式业务操作"],
    ], [3.0, 3.0, 8.8])
    doc.add_heading("4.2 功能详细描述", level=2)
    sections = [
        ("4.2.1 用户管理", [
            ("用户登录", "用户输入账号和密码，后端校验密码哈希后返回访问令牌和用户信息。管理端未登录访问受保护页面时自动跳转到登录页。"),
            ("个人信息管理", "员工可在移动端查看和编辑基本资料，管理员可在后台维护姓名、手机号、角色、项目绑定和在职状态。"),
            ("会话管理", "系统通过 sessions 表记录 refresh token、设备、平台、IP 和过期时间，便于多端登录和后续安全审计。"),
        ]),
        ("4.2.2 打卡管理", [
            ("外勤打卡", "员工选择打卡类型和项目，系统获取定位并拍摄照片，提交后生成打卡记录，记录经纬度、地址、照片、水印防伪码和围栏状态。"),
            ("自定义类型", "管理员可维护 checkin_types，包括上班、下班、实地考察、进度上报、安全检查等分类。"),
            ("异常处理", "围栏外、缺少定位、离线同步等记录进入管理端列表，主管可结合照片和轨迹进行复核。"),
        ]),
        ("4.2.3 位置与轨迹管理", [
            ("位置上报", "移动端定期或在业务节点上报位置，后端保存精度、速度、地址和时间。"),
            ("轨迹回放", "管理端根据人员和日期查询轨迹点，展示总点数、总距离、停留点和由打卡点组成的时间线。"),
            ("热力图", "系统可按日期范围和项目生成热力数据，辅助分析外勤覆盖区域和人员活动密度。"),
        ]),
        ("4.2.4 排班与考勤管理", [
            ("班次维护", "管理员创建早班、晚班、夜班等班次，并设置上下班时间、迟到容忍和早退容忍。"),
            ("批量排班", "主管按人员、日期范围和休息日批量生成排班，移动端可查看今日排班。"),
            ("考勤分析", "系统定时读取打卡记录、排班和考勤组规则，生成正常、迟到、早退、缺勤、请假、休息等结果。"),
        ]),
        ("4.2.5 审批管理", [
            ("提交申请", "员工可发起补卡、请假、加班申请，填写原因、开始日期和结束日期。"),
            ("审批处理", "主管在管理端审批通过或驳回，并填写备注；审批结果联动考勤统计。"),
            ("状态查询", "员工可查看待审批、已通过和已驳回记录，管理端可按状态筛选。"),
        ]),
        ("4.2.6 管理后台", [
            ("工作台", "展示人员、打卡、异常、审批和项目统计，支持图表化查看整体运行情况。"),
            ("基础资料", "提供人员、项目、设备、水印模板、考勤组、班次、节假日等管理页面。"),
            ("数据导出", "通过后端 ExcelJS 生成报表，便于线下归档和进一步统计分析。"),
        ]),
        ("4.2.7 消息与 AI 助手", [
            ("消息通知", "系统可向用户发送考勤、项目、告警和系统类通知，移动端与管理端均可查看。"),
            ("即时通讯", "通过 OpenIM 提供单聊、会话列表和消息同步能力，增强团队沟通。"),
            ("AI 查询", "用户以自然语言询问考勤、审批、项目或报表数据，系统返回结构化结果；执行类操作必须二次确认。"),
        ]),
    ]
    for title, entries in sections:
        doc.add_heading(title, level=3)
        for name, desc in entries:
            p = add_para(doc)
            r = p.add_run(name + "：")
            r.bold = True
            p.add_run(desc)

    doc.add_heading("五、UML系统建模", level=1)
    doc.add_heading("5.1 UML统一建模", level=2)
    add_para(doc, "UML 建模用于从不同视角描述系统需求。本系统以员工、主管/管理员和外部服务为主要参与者，通过用例图描述功能边界，通过序列/流程图描述外勤打卡链路，通过 ER 图描述数据库实体关系。")
    doc.add_heading("5.2 UML用例图", level=2)
    use_case = make_use_case()
    doc.add_picture(str(use_case), width=Cm(15.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("5.3 UML序列图", level=2)
    add_para(doc, "以外勤打卡为例，典型序列为：员工打开移动端打卡页；移动端请求项目、排班和水印模板；移动端调用定位服务获取经纬度和地址；员工拍照确认后提交打卡；后端校验 JWT、围栏距离和打卡类型；后端写入 checkins、防伪码和通知数据；移动端展示打卡结果；管理端可实时查看记录。")
    doc.add_heading("5.4 系统流程图", level=2)
    flow = make_flow()
    doc.add_picture(str(flow), width=Cm(15.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("六、系统的数据库设计", level=1)
    doc.add_heading("6.1 数据库概念设计", level=2)
    add_para(doc, "系统数据库围绕用户、项目、打卡、定位、排班、考勤结果、审批、通知、水印、会话和 AI 交互等实体展开。MySQL 负责强结构化业务数据，MongoDB 作为可选扩展用于非核心或扩展型数据。")
    er = make_er()
    doc.add_picture(str(er), width=Cm(15.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("6.2 关系模型向E-R图的转换", level=2)
    add_bullets(doc, [
        "users 与 checkins、locations、approval_requests、notifications、sessions 为一对多关系。",
        "projects 与 users、checkins、devices、attendance_groups 为关联关系，用于限定外勤地点和项目归属。",
        "attendance_groups 通过 attendance_group_members 与 users 形成多对多关系。",
        "user_schedules 将用户、日期和班次关联起来，attendance_results 保存分析后的每日考勤结论。",
        "ai_conversations、ai_messages、ai_action_requests 和 ai_audit_logs 共同记录 AI 助手的会话、消息、待确认操作和审计结果。"
    ])
    doc.add_heading("6.3 系统流程图", level=2)
    add_para(doc, "系统主要流程包括登录鉴权、移动打卡、离线同步、排班统计、审批处理和消息通知。所有写操作均经过后端接口处理，由控制器完成参数接收，服务层完成业务规则处理，模型层完成数据库访问。")
    doc.add_heading("6.4 数据库逻辑结构设计", level=2)
    db_tables = {
        "users 用户表": [
            ("id", "INT", "是", "用户ID"),
            ("username", "VARCHAR(50)", "否", "登录账号，唯一"),
            ("password", "VARCHAR(255)", "否", "密码哈希"),
            ("name", "VARCHAR(100)", "否", "姓名"),
            ("role", "ENUM", "否", "admin/manager/worker"),
            ("phone", "VARCHAR(20)", "否", "手机号"),
            ("project_id", "INT", "否", "绑定项目"),
            ("status", "TINYINT", "否", "在职状态"),
        ],
        "projects 项目表": [
            ("id", "INT", "是", "项目ID"),
            ("name", "VARCHAR(200)", "否", "项目名称"),
            ("address", "VARCHAR(500)", "否", "项目地址"),
            ("latitude", "DOUBLE", "否", "中心点纬度"),
            ("longitude", "DOUBLE", "否", "中心点经度"),
            ("radius", "INT", "否", "打卡围栏半径"),
            ("status", "TINYINT", "否", "启用状态"),
        ],
        "checkins 打卡记录表": [
            ("id", "INT", "是", "打卡ID"),
            ("user_id", "INT", "否", "用户ID"),
            ("project_id", "INT", "否", "项目ID"),
            ("type", "VARCHAR(50)", "否", "打卡类型"),
            ("latitude/longitude", "DOUBLE", "否", "打卡位置"),
            ("photo", "VARCHAR(500)", "否", "照片路径"),
            ("is_outside", "TINYINT", "否", "是否围栏外"),
            ("distance_to_fence", "DOUBLE", "否", "距离围栏距离"),
            ("watermark_code", "VARCHAR(20)", "否", "水印防伪码"),
        ],
        "offline_checkins 离线打卡表": [
            ("id", "INT", "是", "离线记录ID"),
            ("user_id", "INT", "否", "用户ID"),
            ("type", "VARCHAR(50)", "否", "打卡类型"),
            ("local_timestamp", "DATETIME", "否", "本地发生时间"),
            ("synced", "TINYINT", "否", "同步状态"),
            ("sync_time", "DATETIME", "否", "同步时间"),
        ],
        "attendance_groups 考勤组表": [
            ("id", "INT", "是", "考勤组ID"),
            ("name", "VARCHAR(100)", "否", "考勤组名称"),
            ("work_start_time", "TIME", "否", "上班时间"),
            ("work_end_time", "TIME", "否", "下班时间"),
            ("late_tolerance", "INT", "否", "迟到容忍分钟"),
            ("early_leave_tolerance", "INT", "否", "早退容忍分钟"),
            ("project_id", "INT", "否", "绑定项目"),
        ],
        "approval_requests 审批申请表": [
            ("id", "INT", "是", "审批ID"),
            ("type", "ENUM", "否", "补卡/请假/加班"),
            ("user_id", "INT", "否", "申请人"),
            ("reason", "TEXT", "否", "申请原因"),
            ("start_date/end_date", "DATE", "否", "起止日期"),
            ("status", "ENUM", "否", "pending/approved/rejected"),
            ("approver_id", "INT", "否", "审批人"),
        ],
        "ai_action_requests AI操作确认表": [
            ("id", "INT", "是", "操作ID"),
            ("conversation_id", "INT", "否", "会话ID"),
            ("user_id", "INT", "否", "发起人"),
            ("action_type", "VARCHAR(80)", "否", "操作类型"),
            ("payload", "JSON", "否", "操作参数"),
            ("status", "ENUM", "否", "pending/confirmed/rejected等"),
            ("result", "JSON", "否", "执行结果"),
        ],
    }
    for title, rows in db_tables.items():
        doc.add_heading(title, level=3)
        add_table(doc, ["字段名", "数据类型", "主键", "描述"], rows, [3.5, 3.3, 1.8, 6.2])

    doc.add_heading("七、结束语", level=1)
    add_para(doc, "境图考勤系统围绕企业外勤考勤和项目现场管理需求，完成了从移动端打卡、定位水印、离线同步到后台排班、审批、统计、轨迹回放和 AI 数据助理的需求分析。通过本说明书可以明确系统目标、业务边界、功能模块、数据模型和关键技术方案，为后续系统设计、编码实现、测试部署和课程验收提供依据。")
    add_para(doc, "后续开发中，应继续完善权限细分、异常考勤处理规则、报表导出模板、OpenIM 消息体验和 AI 工具安全策略，并结合实际用户反馈持续优化移动端打卡流程和管理端数据看板。")
    doc.add_heading("参考文献", level=1)
    refs = [
        "[1] Ian Sommerville. Software Engineering. Pearson Education.",
        "[2] Express.js 官方文档：https://expressjs.com/",
        "[3] Flutter 官方文档：https://docs.flutter.dev/",
        "[4] Vue.js 官方文档：https://vuejs.org/",
        "[5] MySQL 8.0 Reference Manual：https://dev.mysql.com/doc/",
        "[6] OpenIM 官方文档：https://docs.openim.io/",
        "[7] 高德开放平台文档：https://lbs.amap.com/",
        "[8] 项目源码：/Users/eva/Desktop/Project/README.md、BackEnd/server/API.md、BackEnd/server/models/db.js。"
    ]
    for ref in refs:
        add_para(doc, ref)

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = "境图考勤系统需求分析说明书"

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
